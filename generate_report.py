from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(214, 40, 40)
DARK = RGBColor(20, 33, 61)
MUTED = RGBColor(92, 103, 125)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    if color:
      run.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = DARK
    run.font.size = Pt(18 if level == 1 else 14)
    paragraph.space_before = Pt(10)
    paragraph.space_after = Pt(6)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.font.size = Pt(10.5)
    run.font.color.rgb = DARK
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        run.font.size = Pt(10.5)
        run.font.color.rgb = DARK


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        set_cell_text(header_cells[index], header, bold=True, color=RGBColor(255, 255, 255))
        shade_cell(header_cells[index], "14213D")

    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            shade_cell(cells[index], "F8FBFF" if len(table.rows) % 2 == 0 else "FFF7E8")

    doc.add_paragraph()


document = Document()
section = document.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

title = document.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(10)
run = title.add_run("AI-Based Smart Complaint Management System")
run.bold = True
run.font.size = Pt(24)
run.font.color.rgb = DARK

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle_run = subtitle.add_run("MERN Stack Examination Project Report")
subtitle_run.italic = True
subtitle_run.font.size = Pt(13)
subtitle_run.font.color.rgb = MUTED

meta_box = document.add_table(rows=3, cols=2)
meta_box.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_box.style = "Table Grid"
meta_fields = [
    ("Course", "AI Driven Full Stack Development (AI308B)"),
    ("Case Study", "AI-Based Smart Complaint Management System"),
    ("Prepared For", "B.Tech 4th Semester ESE 2025-26"),
]
for row_index, (label, value) in enumerate(meta_fields):
    label_cell, value_cell = meta_box.rows[row_index].cells
    set_cell_text(label_cell, label, bold=True, color=RGBColor(255, 255, 255))
    set_cell_text(value_cell, value)
    shade_cell(label_cell, "D62828")
    shade_cell(value_cell, "FFF7E8")

document.add_paragraph()
add_body(
    document,
    "This project implements a full MERN Stack complaint portal where citizens can register and track complaints online. The system uses AI-assisted analysis to classify complaint priority, summarize the issue, recommend the concerned department, and generate an automated response message.",
)

add_heading(document, "1. Objectives")
add_bullets(
    document,
    [
        "Build a secure MERN Stack application for complaint registration and tracking.",
        "Apply AI analysis for urgency detection, department recommendation, summary generation, and automated responses.",
        "Implement JWT authentication, protected routes, and bcrypt password hashing.",
        "Prepare the application for GitHub upload and Render deployment.",
    ],
)

add_heading(document, "2. Frontend Modules")
add_table(
    document,
    ["Module", "Description"],
    [
        ["Complaint Registration Form", "Captures user name, email, title, description, category, location, and status."],
        ["Complaint List Page", "Displays all complaints with category filter, text search, and AI result cards."],
        ["Complaint Status Update Page", "Allows users and admins to update complaint status values."],
        ["AI Analysis Result Display", "Shows priority, department suggestion, summary, auto-response, and AI source."],
    ],
)

add_heading(document, "3. Backend APIs")
add_table(
    document,
    ["Endpoint", "Purpose"],
    [
        ["POST /api/auth/signup", "Register a new user or admin and generate JWT token."],
        ["POST /api/auth/login", "Authenticate user and return signed JWT token."],
        ["GET /api/auth/me", "Fetch logged-in user profile using protected route."],
        ["POST /api/complaints", "Store a new complaint with AI analysis fields."],
        ["GET /api/complaints", "Fetch complaints with filter and search support."],
        ["PUT /api/complaints/:id", "Update complaint status."],
        ["DELETE /api/complaints/:id", "Delete complaint record."],
        ["GET /api/complaints/search?location=Ghaziabad", "Search complaints by location."],
        ["POST /api/ai/analyze", "Generate AI analysis without storing the complaint."],
    ],
)

add_heading(document, "4. MongoDB Schema")
add_body(
    document,
    "The complaint collection stores base complaint details along with AI analysis output. Additional metadata includes creator reference and timestamps for auditability.",
)
add_table(
    document,
    ["Field", "Type / Purpose"],
    [
        ["name, email, title, description, category, location", "Core complaint input fields with validation."],
        ["status", "Pending, In Progress, Resolved, or Rejected."],
        ["priority", "AI-detected urgency level: Low, Medium, High, or Critical."],
        ["department", "Suggested department such as Water Department or Sanitation Department."],
        ["summary", "Short AI-generated summary of the complaint."],
        ["autoResponse", "Automated message shown to the complainant."],
        ["createdBy", "MongoDB ObjectId reference to the authenticated user."],
        ["timestamps", "Automatically stores created and updated dates."],
    ],
)

add_heading(document, "5. Authentication and Security")
add_bullets(
    document,
    [
        "Passwords are hashed using bcrypt before storage.",
        "JWT tokens are generated during signup and login.",
        "Protected middleware blocks access to complaint APIs without valid token.",
        "Input validation is applied using express-validator.",
    ],
)

add_heading(document, "6. AI Integration")
add_body(
    document,
    "The application supports OpenRouter API integration for intelligent analysis. To keep the project functional during exams or network issues, a heuristic fallback engine also classifies complaints and generates meaningful outputs.",
)
add_table(
    document,
    ["Complaint Type", "Expected AI Output"],
    [
        ["Water leakage", "High priority with Water Department recommendation."],
        ["Electricity issue", "High or Critical priority with Electricity Department recommendation."],
        ["Garbage complaint", "Sanitation Department recommendation."],
        ["Long complaint text", "Concise summary and response message."],
    ],
)

add_heading(document, "7. Manual Test Results")
add_table(
    document,
    ["Test Case", "Observed Result"],
    [
        ["User signup", "Signup successful and token generated."],
        ["User login", "Login successful and token generated."],
        ["Add complaint", "Complaint stored successfully."],
        ["AI routing", "Water complaint mapped to Water Department and High priority."],
        ["Update complaint status", "Status updated to Resolved."],
        ["Invalid input", "Validation returned 400 Bad Request."],
    ],
)

add_heading(document, "8. Deployment Readiness")
add_bullets(
    document,
    [
        "Render blueprint provided in render.yaml for frontend and backend services.",
        "Environment templates added for MongoDB URI, JWT secret, client URL, and OpenRouter settings.",
        "Root-level scripts added for development, build, and backend start.",
    ],
)

add_heading(document, "9. Screenshots to Insert Before Submission")
add_bullets(
    document,
    [
        "Signup and login page screenshots",
        "Complaint registration form screenshot",
        "Complaint list and status update screenshot",
        "AI analysis result screenshot",
        "Postman or Thunder Client request screenshots",
        "MongoDB Atlas collection screenshot",
        "Render deployment and live endpoint screenshots",
    ],
)

add_heading(document, "10. Conclusion")
add_body(
    document,
    "The Smart Complaint Management System demonstrates practical MERN integration with AI-assisted automation. It satisfies the examination requirements for frontend modules, backend APIs, MongoDB schema, authentication, AI features, code quality, and deployment preparation.",
)

document.add_section(WD_SECTION.NEW_PAGE)
appendix_title = document.add_paragraph()
appendix_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
appendix_run = appendix_title.add_run("Appendix: Project Files")
appendix_run.bold = True
appendix_run.font.size = Pt(18)
appendix_run.font.color.rgb = DARK

add_bullets(
    document,
    [
        "README.md for setup, endpoints, and deployment guidance",
        "REPORT.md for the editable report source",
        "client/ React frontend source code",
        "server/ Express and MongoDB backend source code",
        "render.yaml for Render deployment",
    ],
)

document.save("AI_Smart_Complaint_Report.docx")
